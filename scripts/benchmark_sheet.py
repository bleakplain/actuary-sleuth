#!/usr/bin/env python3
"""把 confirmation.csv 转成便于人工填写的 xlsx（表头样式、冻结行、下拉框）。"""
from __future__ import annotations

import argparse
import csv
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent))

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

def execute(params: dict) -> dict:
    csv_path = Path(params["csv"])
    xlsx_path = csv_path.with_suffix(".xlsx")
    with csv_path.open(encoding="utf-8-sig") as handle:
        rows = list(csv.reader(handle))
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "确认清单"
    header_fill = PatternFill("solid", fgColor="4472C4")
    for r_idx, row in enumerate(rows, 1):
        for c_idx, value in enumerate(row, 1):
            cell = sheet.cell(row=r_idx, column=c_idx, value=value)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if r_idx == 1:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = header_fill
    for idx, width in enumerate([14, 12, 26, 40, 45, 45, 10, 12, 14][: len(rows[0])], 1):
        sheet.column_dimensions[get_column_letter(idx)].width = width
    sheet.freeze_panes = "A2"
    validation = DataValidation(type="list", formula1='"yes,no"', allow_blank=True)
    sheet.add_data_validation(validation)
    confirmed_col = get_column_letter(rows[0].index("confirmed") + 1)
    validation.add(f"{confirmed_col}2:{confirmed_col}{len(rows)}")
    workbook.save(xlsx_path)
    return {"success": True, "xlsx": str(xlsx_path), "rows": len(rows) - 1}

def main() -> None:
    parser = argparse.ArgumentParser(description="确认清单 CSV 转 xlsx")
    parser.add_argument("--csv", required=True)
    args = parser.parse_args()
    print(__import__("json").dumps(execute(vars(args)), ensure_ascii=False, indent=2))

if __name__ == "__main__":
    main()
