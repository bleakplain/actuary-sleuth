from dataclasses import FrozenInstanceError
from pathlib import Path
from subprocess import CompletedProcess

import pytest

pytest.importorskip("docx")

from docx import Document

from lib.doc_parser import parse_product_document
from lib.doc_parser.models import (
    AuditBlockType,
    AuditDocument,
    Clause,
    DataTable,
    DocumentSection,
    TableType,
)
from lib.doc_parser.pd import docx_parser
from lib.doc_parser.pd import legacy_doc_parser


def _write_docx(path, product_name: str = "") -> None:
    document = Document()
    if product_name:
        document.add_paragraph(product_name)
    table = document.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "条款编号"
    table.rows[0].cells[1].text = "条款内容"
    table.rows[1].cells[0].text = "1.1"
    table.rows[1].cells[1].text = "其他约定\n本条款保留完整原文。"
    document.save(str(path))


def test_unified_blocks_are_frozen_and_cover_every_content_type():
    document = AuditDocument(
        file_name="产品条款.docx",
        file_type=".docx",
        clauses=[Clause("1.1", "其他约定", "普通条款")],
        tables=[DataTable([["年龄", "费率"], ["18", "100"]], TableType.PREMIUM)],
        unclassified_sections=[
            DocumentSection("", "未编号普通原文", "unclassified"),
        ],
        notices=[DocumentSection("投保须知", "须知原文", "notice")],
        health_disclosures=[
            DocumentSection("健康告知", "告知原文", "health_disclosure"),
        ],
        exclusions=[DocumentSection("除外声明", "除外原文", "exclusion")],
        rider_clauses=[Clause("A.1", "附加约定", "附加条款")],
    )

    assert tuple(block.block_type for block in document.audit_blocks) == (
        AuditBlockType.UNCLASSIFIED,
        AuditBlockType.CLAUSE,
        AuditBlockType.TABLE,
        AuditBlockType.NOTICE,
        AuditBlockType.HEALTH_DISCLOSURE,
        AuditBlockType.EXCLUSION,
        AuditBlockType.RIDER,
    )
    assert all(block.document_fingerprint == document.document_fingerprint for block in document.audit_blocks)
    assert len({block.clause_id for block in document.audit_blocks}) == 7
    with pytest.raises(FrozenInstanceError):
        document.audit_blocks[0].content = "不可修改"


def test_fingerprint_and_clause_ids_are_stable_and_filename_independent():
    first = AuditDocument(
        file_name="上传临时名.docx",
        file_type=".docx",
        clauses=[Clause("1.1", "其他约定", "保留原文", topics=())],
    )
    second = AuditDocument(
        file_name="原始产品名.docx",
        file_type=".docx",
        clauses=[Clause("1.1", "其他约定", "保留原文", topics=())],
    )

    assert first.document_fingerprint == second.document_fingerprint
    assert first.audit_blocks[0].clause_id == second.audit_blocks[0].clause_id


def test_clause_id_ignores_derived_topics_and_unrelated_document_changes():
    original = AuditDocument(
        file_name="产品条款.docx",
        file_type=".docx",
        clauses=[
            Clause("1.1", "等待期", "等待期为30日", topics=()),
            Clause("1.2", "其他", "原约定"),
        ],
    )
    changed = AuditDocument(
        file_name="产品条款.docx",
        file_type=".docx",
        clauses=[
            Clause(
                "1.1",
                "等待期",
                "等待期为30日",
                topics=("coverage.waiting_period",),
            ),
            Clause("1.2", "其他", "已修改的无关约定"),
        ],
    )

    assert original.document_fingerprint != changed.document_fingerprint
    assert original.audit_blocks[0].clause_id == changed.audit_blocks[0].clause_id


def test_table_row_and_column_boundaries_change_fingerprint_and_clause_id():
    tabular = AuditDocument(
        file_name="表格.docx",
        file_type=".docx",
        tables=[DataTable(
            [["A", "B"], ["C", "D"]],
            TableType.OTHER,
            raw_text="A\tB\nC\tD",
        )],
    )
    flattened = AuditDocument(
        file_name="表格.docx",
        file_type=".docx",
        tables=[DataTable(
            [["A B C D"]],
            TableType.OTHER,
            raw_text="A B C D",
        )],
    )

    assert tabular.document_fingerprint != flattened.document_fingerprint
    assert tabular.audit_blocks[0].clause_id != flattened.audit_blocks[0].clause_id


def test_audit_document_content_collections_cannot_be_mutated_after_creation():
    document = AuditDocument(
        file_name="产品条款.docx",
        file_type=".docx",
        clauses=[Clause("1.1", "等待期", "等待期为30日")],
    )

    with pytest.raises(AttributeError):
        document.clauses.append(Clause("1.2", "其他", "其他内容"))
    assert len(document.audit_blocks) == 1


def test_unknown_topic_block_is_preserved(tmp_path):
    path = tmp_path / "unknown.docx"
    _write_docx(path)

    result = parse_product_document(str(path))

    clause = next(
        block for block in result.audit_blocks
        if block.block_type is AuditBlockType.CLAUSE
    )
    assert clause.title == "其他约定"
    assert clause.content == "本条款保留完整原文。"
    assert clause.topics == ()
    assert any(
        block.block_type is AuditBlockType.UNCLASSIFIED
        and "条款编号" in block.content
        for block in result.audit_blocks
    )


def test_docx_unnumbered_paragraphs_are_preserved_as_unclassified(tmp_path):
    path = tmp_path / "unnumbered.docx"
    document = Document()
    document.add_paragraph("某某医疗保险条款")
    document.add_paragraph("本段没有条款编号，但仍属于待审核产品原文。")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "1.1"
    table.rows[0].cells[1].text = "保险责任\n本公司承担保险责任。"
    document.save(str(path))

    result = parse_product_document(str(path))

    contents = [
        block.content
        for block in result.audit_blocks
        if block.block_type is AuditBlockType.UNCLASSIFIED
    ]
    combined = "\n".join(contents)
    assert "某某医疗保险条款" in combined
    assert "本段没有条款编号，但仍属于待审核产品原文。" in combined
    assert all(block.topics == () for block in result.audit_blocks
               if block.block_type is AuditBlockType.UNCLASSIFIED)


def test_docx_one_row_table_and_special_heading_are_represented(tmp_path):
    path = tmp_path / "nonstandard.docx"
    document = Document()
    document.add_paragraph("投保须知")
    document.add_paragraph("请投保人完整阅读本须知。")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "特别约定"
    table.rows[0].cells[1].text = "这一行不是标准数据表，但仍须审核。"
    document.save(str(path))

    result = parse_product_document(str(path))

    assert result.notices[0].title == "投保须知"
    assert "请投保人完整阅读本须知。" in result.notices[0].content
    assert "特别约定" in result.notices[0].content
    assert "这一行不是标准数据表，但仍须审核。" in (
        result.notices[0].content
    )


def test_docx_table_remark_uses_the_same_canonical_text_for_product_tags(
    tmp_path,
):
    path = tmp_path / "table-remark.docx"
    document = Document()
    document.add_paragraph("某某医疗保险条款")
    document.add_paragraph("附表一：保证续保期间为六年")
    table = document.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "计划"
    table.rows[0].cells[1].text = "保险金额"
    table.rows[1].cells[0].text = "计划一"
    table.rows[1].cells[1].text = "100万元"
    document.save(str(path))

    result = parse_product_document(str(path))

    assert result.tables[0].remark == "附表一：保证续保期间为六年"
    assert "【数据表 1】附表一：保证续保期间为六年" in result.canonical_text
    assert result.product_tags.renewal_type.value == "guaranteed"


def test_docx_merged_clause_cells_are_not_duplicated_as_unclassified(tmp_path):
    path = tmp_path / "merged-clause.docx"
    document = Document()
    table = document.add_table(rows=2, cols=4)
    table.rows[0].cells[0].text = "条款编号"
    table.rows[0].cells[1].text = "条款标题"
    table.rows[0].cells[2].merge(table.rows[0].cells[3]).text = "条款正文"
    table.rows[1].cells[0].text = "1.1"
    table.rows[1].cells[1].text = "保险责任"
    body = "这一段合并单元格正文只能进入审核输入一次。"
    table.rows[1].cells[2].merge(table.rows[1].cells[3]).text = body
    document.save(str(path))

    result = parse_product_document(str(path))

    assert result.clauses[0].title == "保险责任"
    assert result.clauses[0].text == body
    assert all(
        body not in section.content
        for section in result.unclassified_sections
    )


def test_product_name_priority_and_original_filename_fallback(tmp_path):
    path = tmp_path / "temporary.docx"
    _write_docx(path, "正文安心医疗保险条款")

    user_result = parse_product_document(
        str(path),
        original_file_name="文件名团体医疗保险条款.docx",
        user_product_name="用户指定附加团体医疗保险条款",
    )
    body_result = parse_product_document(
        str(path),
        original_file_name="文件名团体医疗保险条款.docx",
    )

    no_name_path = tmp_path / "temporary-no-name.docx"
    _write_docx(no_name_path)
    file_result = parse_product_document(
        str(no_name_path),
        original_file_name="文件名团体医疗保险条款.docx",
    )
    renamed_result = parse_product_document(
        str(no_name_path),
        original_file_name="另一个上传名称.docx",
    )

    assert user_result.file_name == "文件名团体医疗保险条款.docx"
    assert user_result.product_name == "用户指定附加团体医疗保险条款"
    assert user_result.product_name_source == "user_input"
    assert body_result.product_name == "正文安心医疗保险条款"
    assert body_result.product_name_source == "document_content"
    assert file_result.product_name == "文件名团体医疗保险条款"
    assert file_result.product_name_source == "file_name"
    assert file_result.document_fingerprint == renamed_result.document_fingerprint
    assert tuple(block.clause_id for block in file_result.audit_blocks) == tuple(
        block.clause_id for block in renamed_result.audit_blocks
    )


def test_parser_builds_product_tags_once(tmp_path, monkeypatch):
    path = tmp_path / "temporary.docx"
    _write_docx(path, "正文安心医疗保险条款")
    original_builder = docx_parser.build_product_tags
    generated = []

    def tracked_builder(*args, **kwargs):
        tags = original_builder(*args, **kwargs)
        generated.append(tags)
        return tags

    monkeypatch.setattr(docx_parser, "build_product_tags", tracked_builder)

    result = parse_product_document(
        str(path),
        user_product_name="用户指定医疗保险条款",
    )

    assert len(generated) == 1
    assert result.product_tags is generated[0]


def test_pdf_parser_returns_stable_unified_blocks(tmp_path):
    pytest.importorskip("reportlab")
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas

    path = tmp_path / "sample.pdf"
    pdf = canvas.Canvas(str(path), pagesize=A4)
    pdf.drawString(100, 700, "1.1 Coverage")
    pdf.drawString(100, 680, "Coverage text")
    pdf.save()

    first = parse_product_document(str(path))
    second = parse_product_document(str(path))

    assert first.audit_blocks
    assert first.audit_blocks[0].block_type is AuditBlockType.CLAUSE
    assert first.document_fingerprint == second.document_fingerprint
    assert first.audit_blocks[0].clause_id == second.audit_blocks[0].clause_id


def test_pdf_preamble_is_preserved_once_as_unclassified(tmp_path):
    pytest.importorskip("reportlab")
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas

    path = tmp_path / "preamble.pdf"
    pdf = canvas.Canvas(str(path), pagesize=A4)
    pdf.drawString(100, 720, "Product preamble requiring review")
    pdf.drawString(100, 700, "1.1 Coverage")
    pdf.drawString(100, 680, "Coverage text")
    pdf.save()

    result = parse_product_document(str(path))

    preamble_blocks = [
        block for block in result.audit_blocks
        if block.block_type is AuditBlockType.UNCLASSIFIED
        and "Product preamble requiring review" in block.content
    ]
    assert len(preamble_blocks) == 1
    assert all(
        "Product preamble requiring review" not in clause.text
        for clause in result.clauses
    )


def test_legacy_doc_uses_conversion_and_returns_unified_blocks(tmp_path, monkeypatch):
    source = tmp_path / "旧版医疗保险条款.doc"
    source.write_bytes(
        bytes.fromhex("d0cf11e0a1b11ae1") + b"legacy-placeholder"
    )

    monkeypatch.setattr(legacy_doc_parser.shutil, "which", lambda _: "/fake/soffice")

    def fake_run(command, **kwargs):
        output_dir = command[command.index("--outdir") + 1]
        converted = Path(output_dir) / "旧版医疗保险条款.docx"
        _write_docx(converted, "正文医疗保险条款")
        return CompletedProcess(command, 0, stdout="", stderr="")

    monkeypatch.setattr(legacy_doc_parser.subprocess, "run", fake_run)

    result = parse_product_document(str(source))

    assert result.file_type == ".doc"
    assert result.file_name == source.name
    assert result.audit_blocks
    assert any(
        block.block_type is AuditBlockType.UNCLASSIFIED
        and "正文医疗保险条款" in block.content
        for block in result.audit_blocks
    )
    assert all(block.clause_id for block in result.audit_blocks)
