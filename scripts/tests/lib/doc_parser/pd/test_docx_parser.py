#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Word 解析器测试"""
from copy import deepcopy
import xml.etree.ElementTree as ET
import zipfile

import pytest
from PIL import Image

pytest.importorskip("docx")

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from lib.doc_parser import parse_product_document
from lib.doc_parser.pd.docx_parser import DocxParser


def _add_alternate_content_navigation_image(
    document,
    image_path,
    *,
    hint: str = "请扫描以查询验证条款",
) -> None:
    """构造真实模板使用的 Choice/Fallback 双表示导航图片。"""
    def make_element(namespace: str, local_name: str):
        return document.element.makeelement(f"{{{namespace}}}{local_name}")

    compatibility_namespace = (
        "http://schemas.openxmlformats.org/markup-compatibility/2006"
    )
    shape_namespace = (
        "http://schemas.microsoft.com/office/word/2010/wordprocessingShape"
    )
    vml_namespace = "urn:schemas-microsoft-com:vml"
    paragraph = document.add_paragraph()
    paragraph.add_run().add_picture(str(image_path))
    drawing = next(
        node for node in paragraph._p.iter()
        if node.tag == qn("w:drawing")
    )
    choice_drawing = deepcopy(drawing)
    fallback_drawing = deepcopy(drawing)
    choice_blip = next(
        node for node in choice_drawing.iter()
        if node.tag == qn("a:blip")
    )
    source_relationship_id = choice_blip.get(qn("r:embed"))
    assert source_relationship_id
    source_relationship = document.part.rels[source_relationship_id]
    fallback_relationship_id = "rIdNavigationFallback"
    document.part.rels.add_relationship(
        source_relationship.reltype,
        source_relationship.target_part,
        fallback_relationship_id,
    )
    fallback_blip = next(
        node for node in fallback_drawing.iter()
        if node.tag == qn("a:blip")
    )
    fallback_blip.set(qn("r:embed"), fallback_relationship_id)

    def build_text_box(image_drawing):
        content = OxmlElement("w:txbxContent")
        image_paragraph = OxmlElement("w:p")
        image_run = OxmlElement("w:r")
        image_run.append(image_drawing)
        image_paragraph.append(image_run)
        hint_paragraph = OxmlElement("w:p")
        hint_run = OxmlElement("w:r")
        hint_text = OxmlElement("w:t")
        hint_text.text = hint
        hint_run.append(hint_text)
        hint_paragraph.append(hint_run)
        content.extend((image_paragraph, hint_paragraph))
        return content

    alternate = make_element(compatibility_namespace, "AlternateContent")
    choice = make_element(compatibility_namespace, "Choice")
    choice.set("Requires", "wps")
    choice_shape = make_element(shape_namespace, "wsp")
    choice_text_box = make_element(shape_namespace, "txbx")
    choice_text_box.append(build_text_box(choice_drawing))
    choice_shape.append(choice_text_box)
    choice.append(choice_shape)
    fallback = make_element(compatibility_namespace, "Fallback")
    fallback_picture = OxmlElement("w:pict")
    fallback_shape = make_element(vml_namespace, "rect")
    fallback_text_box = make_element(vml_namespace, "textbox")
    fallback_text_box.append(build_text_box(fallback_drawing))
    fallback_shape.append(fallback_text_box)
    fallback_picture.append(fallback_shape)
    fallback.append(fallback_picture)
    alternate.extend((choice, fallback))
    for child in list(paragraph._p):
        paragraph._p.remove(child)
    run = OxmlElement("w:r")
    run.append(alternate)
    paragraph._p.append(run)


def _add_footnote_part(source, target) -> None:
    content_types_namespace = (
        "http://schemas.openxmlformats.org/package/2006/content-types"
    )
    relationship_namespace = (
        "http://schemas.openxmlformats.org/package/2006/relationships"
    )
    with zipfile.ZipFile(source) as archive:
        entries = {
            name: archive.read(name)
            for name in archive.namelist()
        }

    content_types = ET.fromstring(entries["[Content_Types].xml"])
    ET.SubElement(
        content_types,
        f"{{{content_types_namespace}}}Override",
        {
            "PartName": "/word/footnotes.xml",
            "ContentType": (
                "application/vnd.openxmlformats-officedocument."
                "wordprocessingml.footnotes+xml"
            ),
        },
    )
    entries["[Content_Types].xml"] = ET.tostring(
        content_types,
        encoding="utf-8",
        xml_declaration=True,
    )

    relationships = ET.fromstring(entries["word/_rels/document.xml.rels"])
    ET.SubElement(
        relationships,
        f"{{{relationship_namespace}}}Relationship",
        {
            "Id": "rIdFootnoteTest",
            "Type": (
                "http://schemas.openxmlformats.org/officeDocument/2006/"
                "relationships/footnotes"
            ),
            "Target": "footnotes.xml",
        },
    )
    entries["word/_rels/document.xml.rels"] = ET.tostring(
        relationships,
        encoding="utf-8",
        xml_declaration=True,
    )
    document_xml = ET.fromstring(entries["word/document.xml"])
    paragraphs = [
        node for node in document_xml.iter()
        if node.tag == qn("w:p")
    ]
    run = ET.SubElement(paragraphs[-1], qn("w:r"))
    ET.SubElement(run, qn("w:footnoteReference"), {qn("w:id"): "1"})
    entries["word/document.xml"] = ET.tostring(
        document_xml,
        encoding="utf-8",
        xml_declaration=True,
    )
    entries["word/footnotes.xml"] = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<w:footnotes xmlns:w="http://schemas.openxmlformats.org/'
        'wordprocessingml/2006/main">'
        '<w:footnote w:id="1"><w:p><w:r><w:t>'
        '脚注规定院外购药责任'
        '</w:t></w:r></w:p></w:footnote></w:footnotes>'
    ).encode("utf-8")
    with zipfile.ZipFile(target, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, payload in entries.items():
            archive.writestr(name, payload)


def _add_comment_part(
    source,
    target,
    *,
    start_paragraph_offset: int = -1,
    include_range: bool = True,
) -> None:
    content_types_namespace = (
        "http://schemas.openxmlformats.org/package/2006/content-types"
    )
    relationship_namespace = (
        "http://schemas.openxmlformats.org/package/2006/relationships"
    )
    with zipfile.ZipFile(source) as archive:
        entries = {name: archive.read(name) for name in archive.namelist()}
    content_types = ET.fromstring(entries["[Content_Types].xml"])
    ET.SubElement(
        content_types,
        f"{{{content_types_namespace}}}Override",
        {
            "PartName": "/word/comments.xml",
            "ContentType": (
                "application/vnd.openxmlformats-officedocument."
                "wordprocessingml.comments+xml"
            ),
        },
    )
    entries["[Content_Types].xml"] = ET.tostring(
        content_types, encoding="utf-8", xml_declaration=True,
    )
    relationships = ET.fromstring(entries["word/_rels/document.xml.rels"])
    ET.SubElement(
        relationships,
        f"{{{relationship_namespace}}}Relationship",
        {
            "Id": "rIdCommentTest",
            "Type": (
                "http://schemas.openxmlformats.org/officeDocument/2006/"
                "relationships/comments"
            ),
            "Target": "comments.xml",
        },
    )
    entries["word/_rels/document.xml.rels"] = ET.tostring(
        relationships, encoding="utf-8", xml_declaration=True,
    )
    document_xml = ET.fromstring(entries["word/document.xml"])
    paragraphs = [node for node in document_xml.iter() if node.tag == qn("w:p")]
    if include_range:
        paragraphs[start_paragraph_offset].insert(
            0,
            ET.Element(qn("w:commentRangeStart"), {qn("w:id"): "7"}),
        )
        ET.SubElement(
            paragraphs[-1], qn("w:commentRangeEnd"), {qn("w:id"): "7"},
        )
    marker_run = ET.SubElement(paragraphs[-1], qn("w:r"))
    ET.SubElement(marker_run, qn("w:commentReference"), {qn("w:id"): "7"})
    entries["word/document.xml"] = ET.tostring(
        document_xml, encoding="utf-8", xml_declaration=True,
    )
    entries["word/comments.xml"] = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<w:comments xmlns:w="http://schemas.openxmlformats.org/'
        'wordprocessingml/2006/main">'
        '<w:comment w:id="7" w:author="精算复核员" '
        'w:date="2026-08-13T09:30:00Z"><w:p><w:r><w:t>'
        '请核对等待期措辞；院外购药'
        '</w:t></w:r></w:p></w:comment></w:comments>'
    ).encode("utf-8")
    with zipfile.ZipFile(target, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, payload in entries.items():
            archive.writestr(name, payload)


class TestDocxParser:

    def test_supported_extensions(self):
        assert '.docx' in DocxParser.supported_extensions()

    def test_extract_clauses(self, tmp_path, sample_docx_with_clauses):
        docx_file = tmp_path / "test.docx"
        sample_docx_with_clauses(docx_file, [
            ("1.1", "保险责任", "我们承担以下保险责任..."),
            ("1.2", "责任免除", "因下列情形导致..."),
        ])

        doc = parse_product_document(str(docx_file))
        assert len(doc.clauses) == 2
        assert doc.clauses[0].number == "1.1"
        assert doc.clauses[0].title == "保险责任"

    def test_extract_tables(self, tmp_path, sample_docx_with_premium):
        docx_file = tmp_path / "test.docx"
        sample_docx_with_premium(docx_file)

        doc = parse_product_document(str(docx_file))
        assert len(doc.tables) >= 1

    def test_non_clause_table_filtered(self, tmp_path, sample_docx_with_company_info):
        docx_file = tmp_path / "test.docx"
        sample_docx_with_company_info(docx_file)

        doc = parse_product_document(str(docx_file))
        assert all(c.number not in ['', '公司'] for c in doc.clauses)

    def test_substantive_text_box_is_extracted_into_canonical_text(self, tmp_path):
        source = tmp_path / "text-box.docx"
        document = Document()
        document.add_paragraph("测试医疗保险")
        document.add_paragraph("1 保险责任")
        document.add_paragraph("1.1 医疗保险金")
        document.add_paragraph("本公司承担保险责任。")
        document.add_paragraph("1.2 责任免除")
        document.add_paragraph("法律规定不承担的其他情形。")
        text_box = OxmlElement("w:txbxContent")
        paragraph = OxmlElement("w:p")
        run = OxmlElement("w:r")
        text = OxmlElement("w:t")
        text.text = "2.3 等待期"
        run.append(text)
        paragraph.append(run)
        text_box.append(paragraph)
        body = document.element.body
        body.insert(max(0, len(body) - 1), text_box)
        document.save(source)

        parsed = parse_product_document(str(source))

        assert parsed.coverage_attested
        assert parsed.coverage_attestation.unparsed_text_box_count == 0
        assert "2.3" in parsed.canonical_text
        assert "等待期" in parsed.canonical_text

    def test_embedded_image_disables_full_document_attestation(self, tmp_path):
        source = tmp_path / "image.docx"
        document = Document()
        document.add_paragraph("测试医疗保险")
        document.add_paragraph("1 保险责任")
        document.add_paragraph("本公司承担保险责任。")
        drawing = OxmlElement("w:drawing")
        drawing.append(OxmlElement("a:blip"))
        body = document.element.body
        body.insert(max(0, len(body) - 1), drawing)
        document.save(source)

        parsed = parse_product_document(str(source))

        assert not parsed.coverage_attested
        assert parsed.coverage_attestation.unparsed_image_count == 1
        assert any("嵌入图像" in warning for warning in parsed.warnings)

    def test_body_image_with_qr_finder_patterns_is_ignored_and_counted(
        self,
        tmp_path,
    ):
        source = tmp_path / "qr-image.docx"
        image_path = tmp_path / "qr-like.png"
        image = Image.new("1", (76, 76), 1)
        pixels = image.load()
        assert pixels is not None
        module = 2
        offset = 9
        for module_x, module_y in ((0, 0), (22, 0), (0, 22)):
            left = offset + module_x * module
            top = offset + module_y * module
            for y in range(top, top + 7 * module):
                for x in range(left, left + 7 * module):
                    relative_x = (x - left) // module
                    relative_y = (y - top) // module
                    if (
                        relative_x in {0, 6}
                        or relative_y in {0, 6}
                        or 2 <= relative_x <= 4
                        and 2 <= relative_y <= 4
                    ):
                        pixels[x, y] = 0
        image.save(image_path)
        document = Document()
        document.add_paragraph("测试医疗保险")
        document.add_paragraph("1 保险责任")
        document.add_paragraph("本公司承担保险责任。")
        document.add_picture(str(image_path))
        document.save(source)

        parsed = parse_product_document(str(source))

        assert parsed.coverage_attested
        assert parsed.coverage_attestation.ignored_qr_image_count == 1
        assert parsed.coverage_attestation.unparsed_image_count == 0

    def test_navigation_image_choice_and_fallback_count_as_one_logical_qr(
        self,
        tmp_path,
    ):
        source = tmp_path / "alternate-navigation.docx"
        image_path = tmp_path / "navigation.png"
        Image.new("RGB", (76, 76), "#4472C4").save(image_path)
        document = Document()
        document.add_paragraph("测试医疗保险")
        document.add_paragraph("1 保险责任")
        document.add_paragraph("本公司承担保险责任。")
        _add_alternate_content_navigation_image(document, image_path)
        document.save(source)

        parsed = parse_product_document(str(source))

        assert parsed.coverage_attested
        assert parsed.coverage_attestation.ignored_qr_image_count == 1
        assert parsed.coverage_attestation.unparsed_image_count == 0
        assert (
            parsed.coverage_attestation.ignored_navigation_text_box_count
            == 1
        )

    def test_template_identifier_and_navigation_hint_remain_ignorable(
        self,
        tmp_path,
    ):
        source = tmp_path / "template-identifier-navigation.docx"
        image_path = tmp_path / "navigation.png"
        Image.new("RGB", (76, 76), "#4472C4").save(image_path)
        document = Document()
        document.add_paragraph("测试医疗保险")
        document.add_paragraph("1 保险责任")
        document.add_paragraph("本公司承担保险责任。")
        _add_alternate_content_navigation_image(
            document,
            image_path,
            hint=(
                "人保健康[2025]医疗保险043号\n"
                "请扫描以查询验证条款"
            ),
        )
        document.save(source)

        parsed = parse_product_document(str(source))

        assert parsed.coverage_attested
        assert parsed.coverage_attestation.ignored_qr_image_count == 1
        assert (
            parsed.coverage_attestation.ignored_navigation_text_box_count
            == 1
        )

    def test_navigation_words_do_not_hide_numbered_text_or_exempt_its_image(
        self,
        tmp_path,
    ):
        source = tmp_path / "mixed-navigation-and-clause.docx"
        image_path = tmp_path / "ordinary-navigation-image.png"
        Image.new("RGB", (76, 76), "#4472C4").save(image_path)
        document = Document()
        document.add_paragraph("测试医疗保险")
        document.add_paragraph("1 保险责任")
        document.add_paragraph("本公司承担保险责任。")
        _add_alternate_content_navigation_image(
            document,
            image_path,
            hint=(
                "请扫描以查询验证条款\n"
                "2.3 等待期\n"
                "被保险人等待期内出险不承担保险责任。"
            ),
        )
        document.save(source)

        parsed = parse_product_document(str(source))

        assert not parsed.coverage_attested
        assert parsed.coverage_attestation.ignored_qr_image_count == 0
        assert (
            parsed.coverage_attestation.ignored_navigation_text_box_count
            == 0
        )
        assert parsed.coverage_attestation.unparsed_image_count == 1
        assert "2.3" in parsed.canonical_text
        assert "被保险人等待期内出险" in parsed.canonical_text
        assert any(clause.number == "2.3" for clause in parsed.clauses)

    def test_navigation_hint_outside_image_text_box_does_not_ignore_image(
        self,
        tmp_path,
    ):
        source = tmp_path / "separate-navigation-hint.docx"
        image_path = tmp_path / "ordinary-square.png"
        Image.new("RGB", (76, 76), "#4472C4").save(image_path)
        document = Document()
        document.add_paragraph("测试医疗保险")
        document.add_paragraph("1 保险责任")
        document.add_paragraph("本公司承担保险责任。")
        document.add_picture(str(image_path))
        text_box = OxmlElement("w:txbxContent")
        paragraph = OxmlElement("w:p")
        run = OxmlElement("w:r")
        text = OxmlElement("w:t")
        text.text = "请扫描以查询验证条款"
        run.append(text)
        paragraph.append(run)
        text_box.append(paragraph)
        document.element.body.insert(-1, text_box)
        document.save(source)

        parsed = parse_product_document(str(source))

        assert not parsed.coverage_attested
        assert parsed.coverage_attestation.ignored_qr_image_count == 0
        assert parsed.coverage_attestation.unparsed_image_count == 1

    def test_high_contrast_square_without_qr_finders_remains_unresolved(
        self,
        tmp_path,
    ):
        source = tmp_path / "square-chart.docx"
        image_path = tmp_path / "square-chart.png"
        image = Image.new("1", (76, 76), 1)
        pixels = image.load()
        assert pixels is not None
        for y in range(76):
            for x in range(76):
                if (x // 3 + y // 3) % 2 == 0:
                    pixels[x, y] = 0
        image.save(image_path)
        document = Document()
        document.add_paragraph("测试医疗保险")
        document.add_paragraph("1 保险责任")
        document.add_paragraph("本公司承担保险责任。")
        document.add_picture(str(image_path))
        document.save(source)

        parsed = parse_product_document(str(source))

        assert not parsed.coverage_attested
        assert parsed.coverage_attestation.ignored_qr_image_count == 0
        assert parsed.coverage_attestation.unparsed_image_count == 1

    def test_nonempty_header_and_footer_are_ignored_without_disabling_coverage(
        self,
        tmp_path,
    ):
        source = tmp_path / "header-footer.docx"
        document = Document()
        document.add_paragraph("测试医疗保险")
        document.add_paragraph("1 保险责任")
        document.add_paragraph("本公司承担保险责任。")
        section = document.sections[0]
        section.header.paragraphs[0].text = "页眉中的重大疾病说明"
        section.footer.paragraphs[0].text = "页脚中的药店说明"
        document.save(source)

        parsed = parse_product_document(str(source))

        assert parsed.coverage_attested
        assert (
            parsed.coverage_attestation.ignored_header_footer_part_count
            == 2
        )
        assert "页眉中的重大疾病说明" not in parsed.canonical_text
        assert "页脚中的药店说明" not in parsed.canonical_text
        assert any("页眉/页脚部件" in warning for warning in parsed.warnings)

    def test_empty_header_and_footer_do_not_disable_attestation(self, tmp_path):
        source = tmp_path / "empty-header-footer.docx"
        document = Document()
        document.add_paragraph("测试医疗保险")
        document.add_paragraph("1 保险责任")
        document.add_paragraph("本公司承担保险责任。")
        section = document.sections[0]
        _ = section.header.paragraphs
        _ = section.footer.paragraphs
        document.save(source)

        parsed = parse_product_document(str(source))

        assert parsed.coverage_attested
        assert (
            parsed.coverage_attestation.ignored_header_footer_part_count
            == 0
        )

    def test_header_field_without_literal_text_is_ignored(
        self,
        tmp_path,
    ):
        source = tmp_path / "header-field.docx"
        document = Document()
        document.add_paragraph("测试医疗保险")
        document.add_paragraph("1 保险责任")
        document.add_paragraph("本公司承担保险责任。")
        field = OxmlElement("w:fldSimple")
        field.set(qn("w:instr"), "PAGE")
        document.sections[0].header.paragraphs[0]._p.append(field)
        document.save(source)

        parsed = parse_product_document(str(source))

        assert parsed.coverage_attested
        assert (
            parsed.coverage_attestation.ignored_header_footer_part_count
            == 1
        )

    def test_referenced_footnote_is_extracted_into_canonical_text(
        self,
        tmp_path,
    ):
        base = tmp_path / "base.docx"
        source = tmp_path / "footnote.docx"
        document = Document()
        document.add_paragraph("测试医疗保险")
        document.add_paragraph("1 保险责任")
        document.add_paragraph("1.1 医疗保险金")
        document.add_paragraph("本公司承担保险责任。")
        document.add_paragraph("1.2 责任免除")
        document.add_paragraph("法律规定不承担的其他情形。")
        document.save(base)
        _add_footnote_part(base, source)

        parsed = parse_product_document(str(source))

        assert parsed.coverage_attested
        assert parsed.coverage_attestation.unparsed_auxiliary_part_count == 0
        assert "脚注规定院外购药责任" in parsed.canonical_text
        assert [clause.number for clause in parsed.clauses] == ["1", "1.1", "1.2"]
        assert "【脚注 1】" in parsed.clauses[-1].text
        assert all(
            "【脚注 1】" not in section.content
            for section in parsed.unclassified_sections
        )

    def test_mixed_navigation_prompt_without_image_is_retained(self, tmp_path):
        source = tmp_path / "mixed-navigation-text.docx"
        document = Document()
        document.add_paragraph("测试医疗保险")
        document.add_paragraph("1 保险责任")
        document.add_paragraph("本公司承担保险责任。")
        text_box = OxmlElement("w:txbxContent")
        for value in (
            "扫描二维码查看电子条款",
            "2.3 等待期",
            "等待期为30日。",
        ):
            paragraph = OxmlElement("w:p")
            run = OxmlElement("w:r")
            text = OxmlElement("w:t")
            text.text = value
            run.append(text)
            paragraph.append(run)
            text_box.append(paragraph)
        document.element.body.insert(-1, text_box)
        document.save(source)

        parsed = parse_product_document(str(source))

        assert parsed.coverage_attested
        assert (
            parsed.coverage_attestation.ignored_navigation_text_box_count
            == 0
        )
        waiting_clause = next(
            clause for clause in parsed.clauses if clause.number == "2.3"
        )
        assert "等待期为30日" in waiting_clause.text

    def test_comment_is_retained_but_excluded_from_canonical_text(self, tmp_path):
        base = tmp_path / "base-comment.docx"
        source = tmp_path / "comment.docx"
        document = Document()
        document.add_paragraph("测试医疗保险")
        document.add_paragraph("1 保险责任")
        document.add_paragraph("本公司承担保险责任。")
        document.save(base)
        _add_comment_part(base, source)

        unannotated = parse_product_document(str(base))
        parsed = parse_product_document(str(source))

        assert parsed.coverage_attested
        assert len(parsed.annotations) == 1
        annotation = parsed.annotations[0]
        assert annotation.kind == "comment"
        assert annotation.annotation_id == "7"
        assert annotation.text == "请核对等待期措辞；院外购药"
        assert annotation.author == "精算复核员"
        assert annotation.created_at == "2026-08-13T09:30:00Z"
        assert annotation.anchor_start_paragraph_index == 2
        assert annotation.anchor_end_paragraph_index == 2
        assert annotation.paragraph_index == 2
        assert annotation.text not in parsed.canonical_text
        assert parsed.product_tags.mentions_out_of_hospital_drug is False
        assert parsed.document_fingerprint == unannotated.document_fingerprint
        assert parsed.audit_input_fingerprint == unannotated.audit_input_fingerprint

    def test_comment_range_retains_distinct_paragraph_anchors(self, tmp_path):
        base = tmp_path / "base-spanning-comment.docx"
        source = tmp_path / "spanning-comment.docx"
        document = Document()
        document.add_paragraph("测试医疗保险")
        document.add_paragraph("1 保险责任")
        document.add_paragraph("第一段保险责任。")
        document.add_paragraph("第二段保险责任。")
        document.save(base)
        _add_comment_part(base, source, start_paragraph_offset=-2)

        annotation = parse_product_document(str(source)).annotations[0]

        assert annotation.anchor_start_paragraph_index == 2
        assert annotation.anchor_end_paragraph_index == 3
        assert annotation.paragraph_index == 2

    def test_comment_reference_without_range_is_a_point_anchor(self, tmp_path):
        base = tmp_path / "base-point-comment.docx"
        source = tmp_path / "point-comment.docx"
        document = Document()
        document.add_paragraph("测试医疗保险")
        document.add_paragraph("1 保险责任")
        document.add_paragraph("本公司承担保险责任。")
        document.save(base)
        _add_comment_part(base, source, include_range=False)

        annotation = parse_product_document(str(source)).annotations[0]

        assert annotation.anchor_start_paragraph_index == 2
        assert annotation.anchor_end_paragraph_index == 2
        assert annotation.paragraph_index == 2

    def test_qr_navigation_text_box_is_ignored_without_disabling_coverage(
        self,
        tmp_path,
    ):
        source = tmp_path / "qr-navigation.docx"
        document = Document()
        document.add_paragraph("测试医疗保险")
        document.add_paragraph("1 保险责任")
        document.add_paragraph("本公司承担保险责任。")
        text_box = OxmlElement("w:txbxContent")
        paragraph = OxmlElement("w:p")
        run = OxmlElement("w:r")
        text = OxmlElement("w:t")
        text.text = "扫描二维码查看电子条款"
        run.append(text)
        paragraph.append(run)
        text_box.append(paragraph)
        document.element.body.insert(-1, text_box)
        document.save(source)

        parsed = parse_product_document(str(source))

        assert parsed.coverage_attested
        assert parsed.coverage_attestation.ignored_navigation_text_box_count == 1
        assert "扫描二维码" not in parsed.canonical_text

    def test_template_qr_hint_wording_is_ignored(self, tmp_path):
        source = tmp_path / "template-qr-hint.docx"
        document = Document()
        document.add_paragraph("测试医疗保险")
        document.add_paragraph("1 保险责任")
        document.add_paragraph("本公司承担保险责任。")
        text_box = OxmlElement("w:txbxContent")
        paragraph = OxmlElement("w:p")
        run = OxmlElement("w:r")
        text = OxmlElement("w:t")
        text.text = "请扫描以查询验证条款"
        run.append(text)
        paragraph.append(run)
        text_box.append(paragraph)
        document.element.body.insert(-1, text_box)
        document.save(source)

        parsed = parse_product_document(str(source))

        assert parsed.coverage_attested
        assert parsed.coverage_attestation.ignored_navigation_text_box_count == 1
        assert "请扫描以查询验证条款" not in parsed.canonical_text

    def test_verified_redundant_navigation_text_box_is_ignored(self, tmp_path):
        source = tmp_path / "redundant-navigation.docx"
        document = Document()
        document.add_paragraph("测试医疗保险")
        for number, title in (
            ("1.1", "保险责任"),
            ("1.2", "等待期"),
            ("1.3", "责任免除"),
        ):
            document.add_paragraph(f"{number} {title}")
            document.add_paragraph(f"{title}正文")
        text_box = OxmlElement("w:txbxContent")
        for page, (number, title) in enumerate((
            ("1.1", "保险责任"),
            ("1.2", "等待期"),
            ("1.3", "责任免除"),
        ), start=1):
            paragraph = OxmlElement("w:p")
            run = OxmlElement("w:r")
            text = OxmlElement("w:t")
            text.text = f"{number} {title} {page}"
            run.append(text)
            paragraph.append(run)
            text_box.append(paragraph)
        document.element.body.insert(-1, text_box)
        document.save(source)

        parsed = parse_product_document(str(source))

        assert parsed.coverage_attested
        assert parsed.coverage_attestation.ignored_navigation_text_box_count == 1
        assert parsed.canonical_text.count("保险责任") == 2

    def test_redundant_navigation_with_chapter_headings_is_ignored(
        self,
        tmp_path,
    ):
        source = tmp_path / "chapter-navigation.docx"
        document = Document()
        document.add_paragraph("测试医疗保险")
        for number, title in (
            ("1", "保险责任"),
            ("1.1", "医疗保险金"),
            ("1.2", "等待期"),
            ("1.3", "责任免除"),
        ):
            document.add_paragraph(f"{number} {title}")
            document.add_paragraph(f"{title}正文")
        text_box = OxmlElement("w:txbxContent")
        for value in (
            "1．保险责任",
            "1.1 医疗保险金",
            "1.2 等待期",
            "1.3 责任免除",
        ):
            paragraph = OxmlElement("w:p")
            run = OxmlElement("w:r")
            text = OxmlElement("w:t")
            text.text = value
            run.append(text)
            paragraph.append(run)
            text_box.append(paragraph)
        document.element.body.insert(-1, text_box)
        document.save(source)

        parsed = parse_product_document(str(source))

        assert parsed.coverage_attested
        assert parsed.coverage_attestation.ignored_navigation_text_box_count == 1
        assert [clause.number for clause in parsed.clauses] == [
            "1", "1.1", "1.2", "1.3",
        ]
