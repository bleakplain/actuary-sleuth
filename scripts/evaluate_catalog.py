"""生成法规分类评估 Excel。

对照三个数据源：
1. DOCX 手册 TOC（预期基线）
2. KB references 目录下的 .md 文件
3. LanceDB 中的 chunk 元数据（条款数）

输出 Excel 用于人工评估分类合理性和完整性。
"""
import re
import sys
from pathlib import Path
from collections import defaultdict, Counter

import yaml
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(Path(__file__).parent))
from lib.doc_parser.kb.converter.docx_to_md import (
    _extract_paragraphs, parse_toc, _SECTION_TO_COLLECTION,
)
from lib.rag_engine.registry import RegulationRegistry


DOCX_PATH = r'C:\Users\yuanl\work\actuary-assets\kb\4.产品开发相关法律法规手册2026.4.docx'
REFERENCES_DIR = Path(r'C:\Users\yuanl\work\actuary-assets\kb\references')
LANCEDB_PATH = r'C:\Users\yuanl\work\actuary-assets\kb\v8\lancedb'
OUTPUT_PATH = r'C:\Users\yuanl\work\actuary-sleuth\catalog_evaluation.xlsx'

# 部分 → collection 排序
COLLECTION_ORDER = [
    '00_保险法', '01_负面清单检查', '02_条款费率管理办法',
    '03_健康保险管理办法', '04_普通型人身保险', '05_分红型人身保险',
    '06_短期健康保险', '07_意外伤害保险', '08_互联网保险产品',
    '09_税优健康险', '10_其他监管规定', '11_万能型人身保险',
    '12_其他险种',
]

AUTHORITY_PATTERNS = [
    ('国家金融监督管理总局人身保险监管司', '国家金融监督管理总局人身保险监管司'),
    ('国家金融监督管理总局', '国家金融监督管理总局'),
    ('中国银保监会办公厅', '中国银保监会办公厅'),
    ('中国银保监会人身险部', '中国银保监会人身险部'),
    ('中国银保监会', '中国银保监会'),
    ('中国保监会办公厅', '中国保监会办公厅'),
    ('中国保监会', '中国保监会'),
    ('最高人民法院', '最高人民法院'),
    ('中国银行保险监督管理委员会', '中国银行保险监督管理委员会'),
    ('中国保险监督管理委员会', '中国保险监督管理委员会'),
]

DOC_NUMBER_RE = re.compile(r'[（(]?\s*(金发|保监发|银保监发|银保监办发|银保监规|保监寿险|保监人身险|金规|令)\s*[〔\[]?\s*(\d{4})\s*[〕\]]?\s*[·\-]?\s*(\d+)\s*号?\s*[)）]?')


def parse_authority_and_doc_number(title: str):
    """从法规标题中抽取发文机关 + 文号。"""
    authority = ''
    for pattern, name in AUTHORITY_PATTERNS:
        if title.startswith(pattern) or pattern in title[:30]:
            authority = name
            break
    m = DOC_NUMBER_RE.search(title)
    doc_number = ''
    if m:
        prefix, year, num = m.group(1), m.group(2), m.group(3)
        doc_number = f"{prefix}〔{year}〕{num}号"
    return authority, doc_number


def normalize_title(t: str) -> str:
    """归一化标题用于匹配。"""
    t = re.sub(r'\s+', '', t)
    t = t.replace('（', '(').replace('）', ')')
    t = re.sub(r'[（）()\s]', '', t)
    t = re.sub(r'\d{4}$', '', t)
    return t.lower()


def parse_md_frontmatter(md_path: Path) -> dict:
    """读取 .md frontmatter。"""
    text = md_path.read_text(encoding='utf-8', errors='ignore')
    if not text.startswith('---'):
        return {}
    end = text.find('\n---', 3)
    if end < 0:
        return {}
    try:
        fm = yaml.safe_load(text[3:end])
        return fm or {}
    except Exception:
        return {}


def title_similarity(a: str, b: str) -> float:
    """简单相似度：共同字符比例。"""
    na, nb = normalize_title(a), normalize_title(b)
    if not na or not nb:
        return 0.0
    common = sum(1 for c in na if c in nb)
    return common / max(len(na), len(nb))


def main():
    print('Parsing DOCX TOC...')
    paragraphs = _extract_paragraphs(DOCX_PATH)
    toc_entries = parse_toc(paragraphs)
    print(f'  TOC entries: {len(toc_entries)}')

    print('Scanning references/*.md...')
    md_files = sorted(REFERENCES_DIR.glob('*/' + '*.md'))
    print(f'  md files: {len(md_files)}')

    print('Loading chunks from LanceDB...')
    registry = RegulationRegistry(LANCEDB_PATH)
    chunk_count_by_file = {e.source_file: e.article_count for e in registry.list_all()}
    print(f'  regulations in LanceDB: {len(chunk_count_by_file)}')

    # 按 collection 分组
    toc_by_collection = defaultdict(list)
    for e in toc_entries:
        coll = _SECTION_TO_COLLECTION.get(e.section_name, '10_其他监管规定')
        toc_by_collection[coll].append(e)

    md_by_collection = defaultdict(list)
    for md in md_files:
        coll = md.parent.name
        md_by_collection[coll].append(md)

    # 构建对照行（按 collection 内最相似匹配）
    print('Matching DOCX ↔ md...')
    rows = []
    for coll in COLLECTION_ORDER:
        toc_list = sorted(toc_by_collection.get(coll, []), key=lambda e: e.chapter_num)
        md_list = sorted(md_by_collection.get(coll, []), key=lambda p: p.name)

        # 在 collection 内做贪心匹配：对每个 md 找最相似的 toc
        matched_toc = set()
        matched_md = set()
        pairs = []

        # 第一轮：高相似度匹配
        similarities = []
        for mi, md in enumerate(md_list):
            fm = parse_md_frontmatter(md)
            md_title = fm.get('regulation', '') or md.stem
            for ti, toc in enumerate(toc_list):
                sim = title_similarity(md_title, toc.title)
                if sim >= 0.5:
                    similarities.append((sim, mi, ti))
        similarities.sort(reverse=True)
        for sim, mi, ti in similarities:
            if mi in matched_md or ti in matched_toc:
                continue
            matched_md.add(mi)
            matched_toc.add(ti)
            pairs.append((md_list[mi], toc_list[ti], sim))

        # 第二轮：剩余 md 找剩余 toc
        for mi, md in enumerate(md_list):
            if mi in matched_md:
                continue
            fm = parse_md_frontmatter(md)
            md_title = fm.get('regulation', '') or md.stem
            best_ti = -1
            best_sim = 0
            for ti, toc in enumerate(toc_list):
                if ti in matched_toc:
                    continue
                sim = title_similarity(md_title, toc.title)
                if sim > best_sim:
                    best_sim = sim
                    best_ti = ti
            if best_ti >= 0:
                matched_md.add(mi)
                matched_toc.add(best_ti)
                pairs.append((md_list[mi], toc_list[best_ti], best_sim))

        # 输出三类记录：匹配对、仅 DOCX、仅 md
        for md, toc, sim in pairs:
            fm = parse_md_frontmatter(md)
            md_title = fm.get('regulation', '') or md.stem
            authority, doc_number = parse_authority_and_doc_number(md_title)
            rows.append({
                'collection': coll,
                'docx_chapter': f'第{toc.chapter_num}章' if toc else '',
                'docx_section': toc.section_name if toc else '',
                'docx_title': toc.title if toc else '',
                'md_file': md.name,
                'md_regulation': md_title,
                'authority': authority,
                'doc_number': doc_number,
                'insurance_type': fm.get('险种类型', ''),
                'chunk_count': chunk_count_by_file.get(md.name, 0),
                'similarity': round(sim, 2),
                'status': 'matched' if sim >= 0.5 else ('weak' if toc else 'extra'),
            })
        # 仅 DOCX
        for ti, toc in enumerate(toc_list):
            if ti not in matched_toc:
                rows.append({
                    'collection': coll,
                    'docx_chapter': f'第{toc.chapter_num}章',
                    'docx_section': toc.section_name,
                    'docx_title': toc.title,
                    'md_file': '(缺失)',
                    'md_regulation': '',
                    'authority': '',
                    'doc_number': '',
                    'insurance_type': '',
                    'chunk_count': 0,
                    'similarity': 0.0,
                    'status': 'missing_in_kb',
                })
        # 仅 md（多出的）
        for mi, md in enumerate(md_list):
            if mi not in matched_md:
                fm = parse_md_frontmatter(md)
                md_title = fm.get('regulation', '') or md.stem
                authority, doc_number = parse_authority_and_doc_number(md_title)
                rows.append({
                    'collection': coll,
                    'docx_chapter': '',
                    'docx_section': '',
                    'docx_title': '',
                    'md_file': md.name,
                    'md_regulation': md_title,
                    'authority': authority,
                    'doc_number': doc_number,
                    'insurance_type': fm.get('险种类型', ''),
                    'chunk_count': chunk_count_by_file.get(md.name, 0),
                    'similarity': 0.0,
                    'status': 'extra_in_kb',
                })

    # === 构建 Excel ===
    print('Building Excel...')
    wb = Workbook()

    thin = Side(style='thin', color='BFBFBF')
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    header_font = Font(name='微软雅黑', bold=True, color='FFFFFF')
    header_fill = PatternFill('solid', fgColor='4472C4')
    header_align = Alignment(horizontal='center', vertical='center', wrap_text=True)
    cell_align = Alignment(vertical='center', wrap_text=True)

    status_fills = {
        'matched': PatternFill('solid', fgColor='FFFFFF'),
        'weak': PatternFill('solid', fgColor='FFF2CC'),     # 浅黄
        'missing_in_kb': PatternFill('solid', fgColor='FFC7CE'),  # 红
        'extra_in_kb': PatternFill('solid', fgColor='C6EFCE'),    # 绿
    }

    # === Sheet 1: 总览 ===
    ws = wb.active
    ws.title = '总览'
    ws['A1'] = '法规分类评估总览'
    ws['A1'].font = Font(name='微软雅黑', size=14, bold=True)
    ws.merge_cells('A1:D1')
    ws['A1'].alignment = Alignment(horizontal='center', vertical='center')

    overview = [
        ('数据源', '数量'),
        ('DOCX TOC 章节数', len(toc_entries)),
        ('KB md 文件数', len(md_files)),
        ('LanceDB 法规数', len(chunk_count_by_file)),
        ('', ''),
        ('分类', 'DOCX 章节数', 'KB md 数', '差值'),
    ]
    for coll in COLLECTION_ORDER:
        n_toc = len(toc_by_collection.get(coll, []))
        n_md = len(md_by_collection.get(coll, []))
        overview.append((coll, n_toc, n_md, n_md - n_toc))
    overview.append(('合计',
                     len(toc_entries),
                     len(md_files),
                     len(md_files) - len(toc_entries)))

    for i, row in enumerate(overview, start=2):
        for j, v in enumerate(row, start=1):
            cell = ws.cell(row=i, column=j, value=v)
            cell.border = border
            cell.alignment = cell_align
            if i == 7 or i == len(overview) + 1:
                cell.font = Font(bold=True)

    ws.column_dimensions['A'].width = 32
    for c in 'BCD':
        ws.column_dimensions[c].width = 16

    # === Sheet 2: 完整对照表 ===
    ws = wb.create_sheet('完整对照表')
    headers = [
        '分类(collection)', 'DOCX 章节', 'DOCX 部分', 'DOCX 标题',
        'KB md 文件', 'KB 法规名(regulation)', '发文机关', '文号',
        '险种类型', '条款单元数', '相似度', '状态',
    ]
    for col, h in enumerate(headers, start=1):
        c = ws.cell(row=1, column=col, value=h)
        c.font = header_font
        c.fill = header_fill
        c.alignment = header_align
        c.border = border

    for i, r in enumerate(rows, start=2):
        values = [
            r['collection'], r['docx_chapter'], r['docx_section'], r['docx_title'],
            r['md_file'], r['md_regulation'], r['authority'], r['doc_number'],
            r['insurance_type'], r['chunk_count'], r['similarity'], r['status'],
        ]
        for col, v in enumerate(values, start=1):
            c = ws.cell(row=i, column=col, value=v)
            c.alignment = cell_align
            c.border = border
        fill = status_fills.get(r['status'])
        if fill:
            for col in range(1, len(headers) + 1):
                ws.cell(row=i, column=col).fill = fill

    widths = [22, 11, 11, 45, 50, 50, 16, 18, 14, 11, 9, 14]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = f'A1:L{len(rows)+1}'

    # === Sheet 3: 差异重点 ===
    ws = wb.create_sheet('差异重点')
    diff_headers = ['类型', '分类', 'DOCX 章节', 'DOCX 标题', 'KB md 文件', 'KB 法规名', '相似度', '说明']
    for col, h in enumerate(diff_headers, start=1):
        c = ws.cell(row=1, column=col, value=h)
        c.font = header_font
        c.fill = header_fill
        c.alignment = header_align
        c.border = border

    diff_row_idx = 2
    for r in rows:
        if r['status'] in ('missing_in_kb', 'extra_in_kb', 'weak'):
            if r['status'] == 'missing_in_kb':
                desc = 'DOCX 中存在但 KB 缺失'
            elif r['status'] == 'extra_in_kb':
                desc = 'KB 中存在但 DOCX 没有对应章节'
            else:
                desc = '匹配相似度低，可能错配'
            values = [
                r['status'], r['collection'], r['docx_chapter'], r['docx_title'],
                r['md_file'], r['md_regulation'], r['similarity'], desc,
            ]
            for col, v in enumerate(values, start=1):
                c = ws.cell(row=diff_row_idx, column=col, value=v)
                c.alignment = cell_align
                c.border = border
                c.fill = status_fills.get(r['status'], PatternFill())
            diff_row_idx += 1

    widths = [16, 22, 11, 45, 50, 50, 9, 32]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = 'A2'

    # === Sheet 4: 分类合理性 ===
    ws = wb.create_sheet('分类合理性')
    cat_headers = ['分类', '序号', '法规名', '发文机关', '文号', '条款数', '分类关键词', '人工判断']
    for col, h in enumerate(cat_headers, start=1):
        c = ws.cell(row=1, column=col, value=h)
        c.font = header_font
        c.fill = header_fill
        c.alignment = header_align
        c.border = border

    # 简单的分类关键词（用于人工核对）
    cat_keywords = {
        '00_保险法': '保险法 / 司法解释',
        '01_负面清单检查': '负面清单',
        '02_条款费率管理办法': '条款 / 费率 / 信息披露',
        '03_健康保险管理办法': '健康保险 / 重疾 / 医疗 / 健康管理',
        '04_普通型人身保险': '普通型 / 精算规定 / 费率',
        '05_分红型人身保险': '分红 / 生命表',
        '06_短期健康保险': '短期健康',
        '07_意外伤害保险': '意外伤害',
        '08_互联网保险产品': '互联网',
        '09_税优健康险': '税优 / 个税',
        '10_其他监管规定': '(杂项 - 重点关注)',
        '11_万能型人身保险': '万能型',
        '12_其他险种': '(杂项 - 重点关注)',
    }

    r_idx = 2
    for coll in COLLECTION_ORDER:
        cat_rows = [r for r in rows if r['collection'] == coll
                    and r['status'] in ('matched', 'extra_in_kb', 'weak')]
        if not cat_rows:
            continue
        for i, r in enumerate(cat_rows, start=1):
            values = [
                coll, i, r['md_regulation'], r['authority'], r['doc_number'],
                r['chunk_count'], cat_keywords.get(coll, ''), '',
            ]
            for col, v in enumerate(values, start=1):
                c = ws.cell(row=r_idx, column=col, value=v)
                c.alignment = cell_align
                c.border = border
            r_idx += 1

    widths = [22, 6, 55, 22, 18, 8, 30, 20]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = 'A2'

    wb.save(OUTPUT_PATH)
    print(f'\nSaved: {OUTPUT_PATH}')
    print(f'Size: {Path(OUTPUT_PATH).stat().st_size} bytes')
    print(f'Sheets: {wb.sheetnames}')

    # 控制台总结
    print('\n=== 关键差异 ===')
    status_count = Counter(r['status'] for r in rows)
    for s, n in status_count.items():
        print(f'  {s}: {n}')

    print('\n=== 分类数量差异 ===')
    for coll in COLLECTION_ORDER:
        n_toc = len(toc_by_collection.get(coll, []))
        n_md = len(md_by_collection.get(coll, []))
        diff = n_md - n_toc
        if diff != 0:
            print(f'  {coll}: DOCX {n_toc} → KB {n_md}  ({diff:+d})')


if __name__ == '__main__':
    main()
